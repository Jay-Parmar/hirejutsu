from datetime import date

import pytest
from docx import Document
from fakes import FIXTURES, fixture
from pypdf import PdfWriter

from hirejutsu.config import Config
from hirejutsu.errors import HirejutsuError
from hirejutsu.resume import find_skills, find_titles, find_years, load_profile, read_text, skill_terms, spellings

TODAY = date(2026, 10, 17)


def test_load_profile_reads_text_resume(tmp_path):
    (tmp_path / "resume.txt").write_text("Python developer\nDjango, AWS", encoding="utf-8")
    profile = load_profile(Config({"resume": "resume.txt"}, tmp_path))
    assert "Django" in profile.text
    assert profile.locations == ["remote"]


def test_load_profile_uses_config_values(tmp_path):
    (tmp_path / "cv.md").write_text("# Me", encoding="utf-8")
    config = Config(
        {
            "resume": "cv.md",
            "profile": {"titles": ["Backend Engineer"], "skills": ["Go"], "locations": ["Pune"], "years": 4},
        },
        tmp_path,
    )
    profile = load_profile(config)
    assert profile.titles == ["Backend Engineer"]
    assert profile.skills == ["Go"]
    assert profile.locations == ["Pune"]
    assert profile.years == 4


def test_load_profile_adds_found_skills_after_configured_ones(tmp_path):
    (tmp_path / "resume.txt").write_text("Python, Django and Kafka", encoding="utf-8")
    config = Config({"resume": "resume.txt", "profile": {"skills": ["kafka", "Rust"]}}, tmp_path)
    assert load_profile(config).skills == ["kafka", "Rust", "Python", "Django"]


def test_load_profile_finds_titles_and_years_unless_configured(tmp_path):
    (tmp_path / "resume.txt").write_text("Data Analyst, Acme, 2015 - 2020", encoding="utf-8")
    found = load_profile(Config({"resume": "resume.txt"}, tmp_path))
    assert (found.titles, found.years) == (["Data Analyst"], 5.0)
    profile = {"titles": ["BI Developer"], "years": 3}
    configured = load_profile(Config({"resume": "resume.txt", "profile": profile}, tmp_path))
    assert (configured.titles, configured.years) == (["BI Developer"], 3)


def test_load_profile_from_a_sample_resume(tmp_path):
    (tmp_path / "resume.txt").write_text(fixture("resume_sample.txt"), encoding="utf-8")
    profile = load_profile(Config({"resume": "resume.txt"}, tmp_path))
    assert profile.titles == ["Senior Software Engineer", "Software Engineer", "Software Engineering Intern"]
    assert {"Python", "Go", "AWS", "PostgreSQL", "Django", "Docker", "Kubernetes"} <= set(profile.skills)
    assert profile.years >= 9.6


def test_missing_resume(tmp_path):
    with pytest.raises(HirejutsuError, match="not found"):
        load_profile(Config({"resume": "nope.txt"}, tmp_path))


def test_unsupported_resume_type(tmp_path):
    (tmp_path / "resume.rtf").write_text("x", encoding="utf-8")
    with pytest.raises(HirejutsuError, match=".rtf"):
        load_profile(Config({"resume": "resume.rtf"}, tmp_path))


def test_read_pdf():
    assert read_text(FIXTURES / "resume.pdf").splitlines()[:2] == ["Priya Sharma", "Senior Software Engineer"]


def test_read_pdf_expands_ligatures():
    # The fixture writes the "fl" in Airflow as one ligature glyph, as LaTeX does.
    assert "Airflow" in read_text(FIXTURES / "resume.pdf")


def test_scanned_pdf_asks_for_a_text_resume(tmp_path):
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(tmp_path / "resume.pdf")
    with pytest.raises(HirejutsuError, match="scanned PDF"):
        read_text(tmp_path / "resume.pdf")


def test_read_pdf_that_is_not_a_pdf(tmp_path):
    (tmp_path / "resume.pdf").write_bytes(b"not a pdf")
    with pytest.raises(HirejutsuError, match="Cannot read"):
        read_text(tmp_path / "resume.pdf")


def test_read_docx_keeps_tables_in_document_order(tmp_path):
    document = Document()
    document.add_paragraph("Priya Sharma")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Experience"
    table.cell(1, 0).text = "Senior Software Engineer"
    table.cell(1, 1).text = "Jan 2021 - Present"
    document.add_paragraph("Python, Django")
    document.save(tmp_path / "resume.docx")
    assert read_text(tmp_path / "resume.docx").splitlines() == [
        "Priya Sharma",
        "Experience",
        "Senior Software Engineer | Jan 2021 - Present",
        "Python, Django",
    ]


def test_read_docx_that_is_not_a_docx(tmp_path):
    (tmp_path / "resume.docx").write_bytes(b"old word file")
    with pytest.raises(HirejutsuError, match="Cannot read"):
        read_text(tmp_path / "resume.docx")


def test_find_skills_uses_main_names_in_order_of_mention():
    text = "Built APIs with django and Postgres on AWS.\nPython since 2015."
    assert find_skills(text) == ["Django", "PostgreSQL", "AWS", "Python"]


def test_find_skills_counts_each_skill_once():
    assert find_skills("JS, ECMAScript and JavaScript; k8s or Kubernetes") == ["JavaScript", "Kubernetes"]


def test_find_skills_matches_whole_words_only():
    assert find_skills("pythonic code, javascripts and laws") == []


def test_find_skills_with_symbols_in_names():
    assert find_skills("c++, c# and CI/CD pipelines") == ["C++", "C#", "CI/CD"]


def test_find_skills_reads_js_libraries_as_javascript():
    assert find_skills("APIs in node.js") == ["Node.js", "JavaScript"]


def test_short_and_plain_word_skills_need_their_exact_case():
    plain = "Ready to go the extra mile, react calmly, spring into action. A swift fix, no rust, each node, c and r."
    assert find_skills(plain) == []
    skills = "Go, R, C, Swift, Rust, Spring, React, Node and Excel"
    assert find_skills(skills) == ["Go", "R", "C", "Swift", "Rust", "Spring", "React", "Node.js", "Excel"]


def test_short_skills_are_not_read_inside_joined_words():
    assert find_skills("R&D for the C-suite, Go-live, USB-C, C++ and J. R. Smith") == ["C++"]


def test_exact_case_skills_have_aliases_in_any_case():
    assert find_skills("golang, reactjs and nodejs") == ["Go", "React", "Node.js"]


def test_skills_list_is_broad_and_lists_each_spelling_once():
    terms = skill_terms()
    assert len({name for _, name, _ in terms}) >= 600
    spellings = [term.lower() for term, _, _ in terms]
    assert len(spellings) == len(set(spellings))


def test_find_skills_beyond_software():
    text = "Tally, GST and MS-Excel; Google Ads, SEO and Figma; agile, Scrum and product management"
    assert find_skills(text) == [
        "Tally", "GST", "Excel", "Google Ads", "SEO", "Figma", "Agile", "Scrum", "Product Management",
    ]


def test_find_skills_for_websites_and_online_stores():
    text = "Built WordPress and Shopify stores, then moved to Webflow, Magento and WooCommerce"
    assert find_skills(text) == ["WordPress", "Shopify", "Webflow", "Magento", "WooCommerce"]


@pytest.mark.parametrize(
    ("text", "years"),
    [
        ("Acme | Jan 2019 - Present", 7.8),
        ("Acme 2018 - 2021", 3.0),
        ("Acme, 03/2020 to 06/2022", 2.3),
        ("Acme, 15/03/2020 - 30/06/2021", 1.3),
        ("Acme Sept. 2021 \u2013 Current", 5.2),
        ("Acme, January 2024 till date", 2.8),
        ("Acme, 2025 - 2030", 1.8),
    ],
)
def test_find_years_from_date_ranges(text, years):
    assert find_years(text, TODAY) == years


def test_find_years_counts_overlaps_once_and_skips_gaps():
    assert find_years("Acme, Jan 2020 - Dec 2021\nFreelance, Jun 2021 - Jun 2022", TODAY) == 2.5
    assert find_years("Acme 2015 - 2016\nGlobex 2018 - 2019", TODAY) == 2.0


def test_find_years_skips_education():
    text = (
        "Bachelor of Engineering\n"
        "2014 - 2018\n"
        "Pune University, 2012 - 2014\n"
        "Software Engineer, Acme, 2018 - 2020\n"
    )
    assert find_years(text, TODAY) == 2.0


def test_find_years_without_dates():
    assert find_years("Python developer", TODAY) is None


@pytest.mark.parametrize(
    ("text", "years"),
    [
        ("Backend developer with 5+ years of experience in Python.", 5),
        ("Over 7 years building data platforms.", 7),
        ("Total Experience: 4.5 Years", 4.5),
        ("3 yrs of hands-on industry experience", 3),
        ("8 years' experience, including 3 years of experience with Go", 8),
    ],
)
def test_find_years_from_stated_experience(text, years):
    assert find_years(text, TODAY) == years


def test_find_years_ignores_other_year_counts():
    assert find_years("B.Tech (4 years), 2 years warranty, a 10 year old app", TODAY) is None


def test_find_years_takes_the_larger_of_stated_and_dated():
    assert find_years("5+ years of experience\nAcme, Jan 2015 - Present", TODAY) == 11.8
    assert find_years("12 years of experience\nAcme, Jan 2020 - Present", TODAY) == 12


def test_find_titles_in_a_sample_resume():
    assert find_titles(fixture("resume_sample.txt")) == [
        "Senior Software Engineer",
        "Software Engineer",
        "Software Engineering Intern",
    ]


def test_find_titles_splits_off_companies_and_labels():
    text = "Product Manager at Google, 2019 - 2021\nAcme Corp - Lead Data Scientist\nRole: Senior Accountant"
    assert find_titles(text) == ["Product Manager", "Lead Data Scientist", "Senior Accountant"]


def test_find_titles_skips_sentences_bullets_and_labels():
    text = (
        "Data analyst with a passion for numbers\n"
        "- Lead Developer on the billing team\n"
        "Worked closely with the Engineering Manager and the sales team\n"
        "Developer Tools: Git, Docker"
    )
    assert find_titles(text) == []


def test_spellings_of_a_skill_by_any_name():
    assert ("Postgres", False) in spellings("postgresql")
    assert spellings("Golang") == [("Go", True), ("Golang", False)]


def test_spellings_of_an_unknown_skill():
    assert spellings("Frobnicate") == [("Frobnicate", False)]
